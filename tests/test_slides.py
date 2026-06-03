"""
Test cho chức năng render slide.

Các test này không gọi Gemini. Mục tiêu là kiểm tra luồng render nội bộ:
SlidePresentation -> slide_builder.js -> file .pptx.
"""

import os
import shutil

import pytest
from pptx import Presentation

from backend.models.schemas import SlideContent, SlidePresentation
from backend.slides.slide_generator import create_pptx_file


def test_create_pptx_file_renders_valid_presentation():
    if shutil.which("node") is None:
        pytest.skip("Cần cài Node.js để render slide")

    presentation_data = SlidePresentation(
        slides=[
            SlideContent(
                title="Mục tiêu bài học",
                bullet_points=[
                    "Hiểu yêu cầu của đồ án",
                    "Nắm được các tiêu chí đánh giá",
                    "Biết cách trình bày kết quả",
                ],
                speaker_notes="Giảng viên giới thiệu tổng quan mục tiêu và cách đánh giá.",
            ),
            SlideContent(
                title="Tiêu chí đánh giá",
                bullet_points=[
                    "Đọc hiểu tài liệu",
                    "Báo cáo kỹ thuật",
                    "Demo thực nghiệm",
                ],
                speaker_notes="Giảng viên giải thích từng tiêu chí và trọng số điểm.",
            ),
        ]
    )

    output_path = create_pptx_file(
        presentation_data=presentation_data,
        output_filename="test_generated_slides.pptx",
        theme="academic",
        author="E-Learning AI",
        subtitle="Slide render test",
    )

    assert os.path.exists(output_path)
    assert os.path.getsize(output_path) > 5000

    pptx = Presentation(output_path)

    # Renderer tạo thêm 1 slide tiêu đề và 1 slide kết thúc.
    assert len(pptx.slides) == len(presentation_data.slides) + 2