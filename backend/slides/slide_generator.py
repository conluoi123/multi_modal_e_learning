import os 
from pptx import Presentation 
from backend.rag.generator import get_llm
from backend.models.schemas import SlidePresentation


from langchain_core.output_parsers import JsonOutputParser
from langchain.prompts import PromptTemplate

def generate_slide_outline(topic: str, context: str, n_slides: int) -> SlidePresentation: 
    """
        AI trả về json để fill vào .pptx
    """
    llm = get_llm()
    parser = JsonOutputParser(pydantic_object=SlidePresentation)
    
    template = """
        Bạn là giáo sư đại học. Hãy tạo bài thuyết trình {n_slides} slides về chủ đề: '{topic}'.
        Sử dụng TÀI LIỆU tham khảo sau:
        {context}
        
        Yêu cầu mỗi slide: Tiêu đề ngắn, 3-4 ý gạch đầu dòng, và có ghi chú (speaker_notes).
        
        {format_instructions}
    """
    prompt = PromptTemplate(
        template=template,
        input_variables=["topic", "context", "n_slides"],
        partial_variables={"format_instructions": parser.get_format_instructions()}
    )
    
    chain = prompt | llm | parser
    
    print(f"AI đang tạo dàn ý cho {n_slides} slides")
    result_dict = chain.invoke({"topic": topic, "context": context, "n_slides": n_slides})
    
    # Parse dict về Pydantic object
    return SlidePresentation(**result_dict)
    
def create_pptx_file(presentation_data: SlidePresentation, output_filename: str): 
    """
        Dùng python-pptx để vẽ PP
    """
    prs = Presentation()
    for slide_data in presentation_data.slides: 
        slide_layout = prs.slide_layouts[1]  # layout tiêu đề + bullet
        slide = prs.slides.add_slide(slide_layout)
        
        # Tiêu đề
        title = slide.shapes.title
        title.text = slide_data.title 

        # Bullet points
        body_shape = slide.placeholders[1]
        tf = body_shape.text_frame 
        for i, bullet in enumerate(slide_data.bullet_points):
            p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
            p.text = bullet 

        notes_slide = slide.notes_slide
        notes_slide.notes_text_frame.text = slide_data.speaker_notes 
    os.makedirs("data/slides", exist_ok=True)
    file_path = f"data/slides/{output_filename}"
    prs.save(file_path)
    print(f" Đã xuất file PowerPoint thành công tại: {file_path}")
    return file_path            

if __name__ == "__main__":
    from backend.rag.retriever import retrive_context
    
    topic = "Tiêu chí đánh giá môn học"
    print("1. Đang lục tìm tài liệu...")
    chunks = retrive_context(topic, k=3)
    
    # Nối text lại
    context_text = "\n".join([c["text"] for c in chunks])
    
    print("2. Bắt đầu sinh cấu trúc Slide...")
    slide_data = generate_slide_outline(topic, context_text, n_slides=3)
    
    print("3. Bắt đầu vẽ file PowerPoint...")
    create_pptx_file(slide_data, "Demo_Bai_Giang.pptx")